"""用虚构的公司导出表结构复现表头截断，不依赖真实业务数据。"""

import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from pptx import Presentation
from PyQt6.QtWidgets import QApplication, QMessageBox

from excel_visualizer.data_service import read_sheet
from excel_visualizer.main_window import MainWindow
from excel_visualizer.ppt_service import export_pptx


def exported_sheet(path, *, merged=False, unnamed=False, far_style=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "虚构问题表"
    sheet.append(["模拟导出说明", "统计记录"])
    if merged:
        sheet.merge_cells("B1:H1")
    sheet.append(["说明：下面才是字段名"])
    sheet.append(["编号", "其他字段", "山头", "创建时间", "说明列", "产品序列号/机台编码", "地区", "服务请求状态"])
    sheet.append([1, "示例", "XRF", "2026-1-2", "虚构", "001", "示例", "已结束"])
    sheet.append([2, "示例", "XRF", "2026-1-3", "虚构", "001", "示例", "处理中"])
    if unnamed:
        sheet["H3"] = None
    if far_style:
        sheet["XFD1000"].fill = PatternFill("solid", fgColor="FFFF00")
    book.save(path)
    book.close()


class ExcelHeaderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "export.xlsx"

    def test_title_row_with_two_cells_must_not_truncate_eight_data_columns(self):
        exported_sheet(self.path)
        sheet = read_sheet(self.path, "虚构问题表", 1)
        self.assertEqual(len(sheet.columns), 8)
        self.assertEqual(sheet.columns[-1], "未命名 [H]")
        self.assertEqual(sheet.rows[-1][-1].value, "处理中")

    def test_merged_title_cells_can_be_shown_as_unnamed_columns(self):
        exported_sheet(self.path, merged=True)
        sheet = read_sheet(self.path, "虚构问题表", 1)
        self.assertEqual(len(sheet.columns), 8)
        self.assertEqual(sheet.columns[2], "未命名 [C]")

    def test_empty_trailing_header_with_real_data_is_preserved(self):
        exported_sheet(self.path, unnamed=True)
        sheet = read_sheet(self.path, "虚构问题表", 3)
        self.assertEqual(len(sheet.columns), 8)
        self.assertEqual(sheet.columns[-1], "未命名 [H]")
        self.assertEqual(sheet.rows[0][-1].value, "已结束")

    def test_merged_cells_inside_header_and_columns_beyond_z(self):
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "名称"
        sheet["B1"] = "合并说明"
        sheet.merge_cells("B1:C1")
        sheet["AA1"] = "末列"
        sheet["AA2"] = "真实内容"
        book.save(self.path)
        book.close()
        data = read_sheet(self.path, "Sheet")
        self.assertEqual(len(data.columns), 27)
        self.assertEqual(data.columns[2], "未命名 [C]")
        self.assertEqual(data.columns[-1], "末列 [AA]")
        self.assertEqual(data.rows[0][-1].value, "真实内容")

    def test_formatting_only_cells_do_not_create_spurious_rows_or_columns(self):
        exported_sheet(self.path, far_style=True)
        sheet = read_sheet(self.path, "虚构问题表", 3)
        self.assertEqual(len(sheet.columns), 8)
        self.assertEqual(sheet.row_numbers, (4, 5))

    def test_gui_header_correction_then_four_page_export(self):
        exported_sheet(self.path, merged=True)
        window = MainWindow()
        self.addCleanup(window.close)
        page = window.issues_page
        with patch.object(QMessageBox, "warning") as warning:
            page.open_excel(self.path)
            self.assertEqual(page.column_boxes["status"].count(), 9)
            self.assertIn("第 3 行", page.status_label.text())
            page.header_row.setValue(3)
            page._load_sheet()
            self.assertEqual({name: box.currentData() for name, box in page.column_boxes.items()},
                             {"product": 2, "created": 3, "machine": 5, "status": 7})
            page._generate_report()
        self.assertFalse(warning.called, str(warning.call_args))
        xrf = page.report.metrics[0]
        self.assertEqual((xrf.total, xrf.machines, xrf.closed, xrf.density, xrf.closure_rate), (2, 1, 1, 2, .5))
        book = Workbook()
        sheet = book.active
        sheet.append(["客户名字", "山头", "uptime", "跑货量"])
        sheet.append(["模拟客户", "XRF", 98, 100])
        monthly_path = self.path.parent / "monthly.xlsx"
        book.save(monthly_path)
        book.close()
        window.open_excel(monthly_path)
        window._generate_report()
        result = export_pptx(window.report, self.path.parent / "report.pptx", issues_report=page.report)
        ppt = Presentation(result)
        self.assertEqual(len(ppt.slides), 4)
        charts = [shape.chart for shape in ppt.slides[3].shapes if shape.has_chart]
        self.assertEqual(len(charts), 4)
        self.assertEqual(charts[0].series[0].values[0], 2)
        self.assertEqual(charts[2].series[0].values[0], 2)
        self.assertEqual(charts[3].series[0].values[0], .5)
