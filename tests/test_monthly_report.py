"""模拟工作簿验证分类、数值转换、界面状态和 PPT 内嵌数据。"""

import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from dataclasses import replace
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from lxml import etree
from matplotlib.colors import to_hex
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Color, PatternFill
from openpyxl.writer.theme import theme_xml
from pptx import Presentation
from PyQt6.QtWidgets import QApplication, QMessageBox
from PyQt6.QtGui import QColor

from excel_visualizer.chart_style import ChartColors, DEFAULT_COLORS
from excel_visualizer.data_service import (
    ExcelDataError, GROUP_TITLES, NON_OPTICAL, OPTICAL, build_report,
    guess_columns, is_yellow, list_sheets, read_sheet,
)
from excel_visualizer.main_window import MainWindow
from excel_visualizer.ppt_service import export_pptx


YELLOW = PatternFill("solid", fgColor="FFFFFF00")
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def fixture(path: Path):
    book = Workbook()
    book.loaded_theme = theme_xml.replace('val="8064A2"', 'val="FFC000"').encode("utf-8")
    book.active.title = "说明"
    sheet = book.create_sheet("月报")
    sheet.append(["模拟月报，仅用于程序验证"])
    sheet.append(["客户名字", "山头", "uptime", "跑货量"])
    for i, product in enumerate(NON_OPTICAL + OPTICAL):
        sheet.append([f"客户{i + 1}", product, 90 + i, 1000 + i * 100])
    sheet["C3"] = 0.98
    sheet["C3"].number_format = "0.0%"
    sheet["C4"] = "97.5%"
    sheet["D4"] = "1,200"
    sheet.append(["未保客户甲", "XRF", 88, 300])
    sheet["A13"].fill = YELLOW
    sheet.append(["未保客户乙", "BFI", 0, 0])
    sheet["A14"].fill = PatternFill("solid", fgColor=Color(indexed=5))
    sheet.append(["未保客户丙", "其他产品", 95, 500])
    sheet["A15"].fill = PatternFill("solid", fgColor=Color(theme=7, tint=0.6))
    book.save(path)
    book.close()


class MonthlyReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "monthly.xlsx"
        fixture(self.path)

    def read(self):
        return read_sheet(self.path, "月报", 2)

    def report(self):
        sheet = self.read()
        return build_report(sheet, guess_columns(sheet))

    def edit(self, operation):
        book = load_workbook(self.path)
        operation(book["月报"])
        book.save(self.path)
        book.close()

    def test_partition_colors_and_percentage_formats(self):
        self.assertEqual(list_sheets(self.path), ["说明", "月报"])
        report = self.report()
        self.assertEqual([len(group.records) for group in report.groups], [4, 6, 3])
        self.assertEqual([r.product for r in report.groups[0].records], list(NON_OPTICAL))
        self.assertEqual([r.product for r in report.groups[1].records], list(OPTICAL))
        self.assertEqual([r.uptime for r in report.groups[0].records[:2]], [98, 97.5])
        self.assertEqual(report.groups[0].records[1].volume, 1200)
        self.assertIn("其他产品", [r.product for r in report.groups[2].records])
        rows = [r.source_row for group in report.groups for r in group.records]
        self.assertEqual(sorted(rows), list(range(3, 16)))

    def test_yellow_range(self):
        for rgb in ("FFFF00", "FFF2CC", "FFD966", "FFFF99"):
            self.assertTrue(is_yellow(rgb), rgb)
        for rgb in (None, "FFFFFF", "00FF00", "FF0000", "0000FF", "FFA500"):
            self.assertFalse(is_yellow(rgb), rgb)

    def test_case_normalization(self):
        self.edit(lambda sheet: setattr(sheet["B3"], "value", " xrf "))
        self.assertEqual(self.report().groups[0].records[0].product, "XRF")

    def test_merged_customers_inherit_yellow(self):
        def update(sheet):
            sheet.merge_cells("A13:A14")
        self.edit(update)
        records = self.report().groups[2].records
        self.assertEqual([r.customer for r in records[:2]], ["未保客户甲", "未保客户甲"])

    def test_yellow_on_metric_does_not_change_warranty(self):
        self.edit(lambda sheet: setattr(sheet["C3"], "fill", YELLOW))
        self.assertEqual(len(self.report().groups[0].records), 4)

    def test_unknown_unmarked_product_is_not_silently_dropped(self):
        self.edit(lambda sheet: setattr(sheet["B3"], "value", "unknown"))
        with self.assertRaisesRegex(ExcelDataError, "第 3 行.*不在指定"):
            self.report()

    def test_invalid_values_block_report(self):
        for cell, value in (("C3", 101), ("C3", -1), ("D3", -3), ("C3", None),
                            ("D3", "no"), ("D3", "12,34"), ("C3", True), ("D3", "NaN")):
            with self.subTest(cell=cell, value=value):
                fixture(self.path)
                def update(sheet):
                    sheet[cell] = value
                    sheet[cell].number_format = "General"
                self.edit(update)
                with self.assertRaisesRegex(ExcelDataError, "第 3 行"):
                    self.report()

    def test_missing_formula_cache_is_explicit(self):
        self.edit(lambda sheet: setattr(sheet["C3"], "value", "=98/100"))
        with self.assertRaisesRegex(ExcelDataError, "公式没有缓存结果"):
            self.report()

    def test_conditional_customer_fill_blocks_ambiguous_classification(self):
        self.edit(lambda sheet: sheet.conditional_formatting.add(
            "A3:A15", CellIsRule(operator="equal", formula=['"客户1"'], fill=YELLOW)))
        with self.assertRaisesRegex(ExcelDataError, "条件格式"):
            self.report()

    def test_duplicate_rows_are_preserved(self):
        self.edit(lambda sheet: sheet.append(["客户2", "XPS", 90, 200]))
        report = self.report()
        self.assertEqual(len(report.groups[0].records), 5)
        self.assertTrue(report.notes)

    def test_unformatted_fraction_requires_explicit_mode(self):
        def update(sheet):
            for row in range(3, 16):
                sheet.cell(row, 3, 0.98).number_format = "General"
        self.edit(update)
        sheet = self.read()
        default = build_report(sheet, guess_columns(sheet))
        fractions = build_report(sheet, guess_columns(sheet), "fraction")
        self.assertEqual(default.groups[0].records[0].uptime, 0.98)
        self.assertEqual(fractions.groups[0].records[0].uptime, 98)

    def test_duplicate_mapping_rejected(self):
        sheet = self.read()
        mapping = guess_columns(sheet)
        mapping["volume"] = mapping["uptime"]
        with self.assertRaisesRegex(ExcelDataError, "不同的列"):
            build_report(sheet, mapping)

    def test_ppt_has_three_editable_combo_charts_and_matching_embedded_values(self):
        report = self.report()
        path = export_pptx(report, Path(self.temp.name) / "report.pptx", "2026-08")
        presentation = Presentation(path)
        self.assertEqual(len(presentation.slides), 3)
        with ZipFile(path) as archive:
            for i, (group, slide) in enumerate(zip(report.groups, presentation.slides), 1):
                charts = [shape.chart for shape in slide.shapes if shape.has_chart]
                self.assertEqual(len(charts), 1)
                self.assertEqual(len(charts[0].plots), 2)
                root = etree.fromstring(archive.read(f"ppt/charts/chart{i}.xml"))
                bar, = root.xpath("//c:barChart", namespaces=NS)
                line, = root.xpath("//c:lineChart", namespaces=NS)
                self.assertEqual(len(bar.xpath("c:ser", namespaces=NS)), 1)
                self.assertEqual(len(line.xpath("c:ser", namespaces=NS)), 1)
                axes = {axis.xpath("c:axId/@val", namespaces=NS)[0]: axis
                        for axis in root.xpath("//c:valAx", namespaces=NS)}
                line_value_id = line.xpath("c:axId/@val", namespaces=NS)[1]
                bar_value_id = bar.xpath("c:axId/@val", namespaces=NS)[1]
                self.assertNotEqual(line_value_id, bar_value_id)
                self.assertEqual(axes[line_value_id].xpath("c:axPos/@val", namespaces=NS), ["l"])
                self.assertEqual(axes[bar_value_id].xpath("c:axPos/@val", namespaces=NS), ["r"])
                self.assertEqual([float(v) for v in line.xpath("c:ser/c:val/c:numRef/c:numCache/c:pt/c:v/text()", namespaces=NS)],
                                 [r.uptime / 100 for r in group.records])
                self.assertEqual([float(v) for v in bar.xpath("c:ser/c:val/c:numRef/c:numCache/c:pt/c:v/text()", namespaces=NS)],
                                 [r.volume for r in group.records])
                embedded = archive.read(f"ppt/embeddings/Microsoft_Excel_Sheet{i}.xlsx")
                book = load_workbook(BytesIO(embedded), data_only=True)
                values = list(book.active.values)
                self.assertEqual([row[0] for row in values[1:]], [r.label for r in group.records])
                for row, record in zip(values[1:], group.records):
                    self.assertAlmostEqual(row[1], record.uptime / 100)
                    self.assertEqual(row[2], record.volume)
                book.close()

    def test_empty_groups_keep_three_slides(self):
        report = self.report()
        report = replace(report, groups=(report.groups[0], replace(report.groups[1], records=()),
                                         replace(report.groups[2], records=())))
        path = export_pptx(report, Path(self.temp.name) / "empty.pptx",
                           colors=ChartColors(line="#112233", bar="#445566"))
        presentation = Presentation(path)
        self.assertEqual(len(presentation.slides), 3)
        for slide in list(presentation.slides)[1:]:
            self.assertIn("本月无符合条件的数据", "".join(s.text for s in slide.shapes if s.has_text_frame))

    def test_failed_export_preserves_existing_file(self):
        path = Path(self.temp.name) / "existing.pptx"
        path.write_bytes(b"existing report")
        with patch("excel_visualizer.ppt_service.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                export_pptx(self.report(), path)
        self.assertEqual(path.read_bytes(), b"existing report")
        self.assertEqual(sorted(p.name for p in Path(self.temp.name).glob("*.pptx")), ["existing.pptx"])

    def test_gui_generate_reset_and_export(self):
        warning = patch.object(QMessageBox, "warning").start()
        self.addCleanup(patch.stopall)
        window = MainWindow()
        self.addCleanup(window.close)
        window.header_row.setValue(2)
        # 示例第一张是说明页，用户可以选择真实数据工作表。
        with patch.object(QMessageBox, "warning"):
            window.open_excel(self.path)
        window.sheet_box.setCurrentText("月报")
        window._load_sheet()
        window._generate_report()
        self.assertIsNotNone(window.report, str(warning.call_args))
        self.assertTrue(window.export_button.isEnabled())
        self.assertEqual(window.table.rowCount(), 13)
        for chart in window.charts:
            chart.draw()
            self.assertEqual(len(chart.figure.axes), 2)
        # 重复生成不得积累双轴。
        window._generate_report()
        self.assertTrue(all(len(chart.figure.axes) == 2 for chart in window.charts))
        output = Path(self.temp.name) / "gui.pptx"
        with patch("excel_visualizer.main_window.QFileDialog.getSaveFileName", return_value=(str(output), "")), \
             patch.object(QMessageBox, "information"):
            window._export_report()
        self.assertEqual(len(Presentation(output).slides), 3)
        window.column_boxes["volume"].setCurrentIndex(1)
        self.assertFalse(window.export_button.isEnabled())
        self.assertIsNone(window.report)

    def test_colors_flow_from_picker_to_preview_and_editable_ppt(self):
        window = MainWindow()
        self.addCleanup(window.close)
        # 导入前可选色，取消另一种颜色不会改变已选颜色。
        with patch("excel_visualizer.main_window.QColorDialog.getColor", return_value=QColor("#208040")):
            window.bar_color_button.click()
        with patch("excel_visualizer.main_window.QColorDialog.getColor", return_value=QColor()):
            window.line_color_button.click()
        self.assertEqual(window.chart_colors, ChartColors(bar="#208040"))
        self.assertFalse(window.export_button.isEnabled())
        window.header_row.setValue(2)
        with patch.object(QMessageBox, "warning"):
            window.open_excel(self.path)
        window.sheet_box.setCurrentText("月报")
        window._load_sheet()
        window._generate_report()
        report = window.report
        self.assertIsNotNone(report)
        window.tabs.setCurrentIndex(3)
        # 生成后改色即时刷新，不重新计算数据或跳离核对页。
        with patch("excel_visualizer.main_window.QColorDialog.getColor", return_value=QColor("#7030a0")):
            window.line_color_button.click()
        self.assertEqual(window.chart_colors, ChartColors(line="#7030A0", bar="#208040"))
        self.assertIs(window.report, report)
        self.assertEqual(window.tabs.currentIndex(), 3)
        self.assertTrue(window.export_button.isEnabled())
        for chart in window.charts:
            chart.draw()
            self.assertEqual(len(chart.figure.axes), 2)
            self.assertEqual(to_hex(chart.axes.lines[0].get_color()).upper(), "#7030A0")
            for bar in chart.figure.axes[1].patches:
                self.assertEqual(to_hex(bar.get_facecolor()).upper(), "#208040")
                self.assertEqual(bar.get_facecolor()[3], 1)
        path = Path(self.temp.name) / "custom-colors.pptx"
        with patch("excel_visualizer.main_window.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
             patch.object(QMessageBox, "information"):
            window.export_button.click()
        self.assertEqual(len(Presentation(path).slides), 3)
        with ZipFile(path) as archive:
            for index in range(1, 4):
                root = etree.fromstring(archive.read(f"ppt/charts/chart{index}.xml"))
                for xpath in (
                    "//c:lineChart/c:ser/c:spPr/a:ln/a:solidFill/a:srgbClr/@val",
                    "//c:lineChart/c:ser/c:marker/c:spPr/a:solidFill/a:srgbClr/@val",
                    "//c:lineChart/c:ser/c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr/@val",
                ):
                    self.assertEqual(root.xpath(xpath, namespaces=NS), ["7030A0"])
                self.assertEqual(root.xpath("//c:barChart/c:ser/c:spPr/a:solidFill/a:srgbClr/@val",
                                            namespaces=NS), ["208040"])
        window.reset_colors_button.click()
        self.assertEqual(window.chart_colors, DEFAULT_COLORS)
        self.assertFalse(window.reset_colors_button.isEnabled())
        for chart in window.charts:
            self.assertEqual(to_hex(chart.axes.lines[0].get_color()).upper(), DEFAULT_COLORS.line)
            self.assertEqual(to_hex(chart.figure.axes[1].patches[0].get_facecolor()).upper(), DEFAULT_COLORS.bar)
        # 切到第二、三步不丢失第一步，第三步仍为空白预留。
        self.assertEqual(window.workflow_tabs.count(), 3)
        for index in (1, 2):
            window.workflow_tabs.setCurrentIndex(index)
            if index == 2:
                self.assertEqual(window.workflow_tabs.currentWidget().children(), [])
            else:
                self.assertIsNone(window.issues_page.report)
        window.workflow_tabs.setCurrentIndex(0)
        self.assertIs(window.report, report)
        self.assertEqual(window.table.rowCount(), 13)
        self.assertTrue(window.export_button.isEnabled())
        # 改数据配置后，换颜色不能意外恢复过期报表。
        window.uptime_mode.setCurrentIndex(1)
        window.set_chart_colors(ChartColors(line="#123456"))
        self.assertIsNone(window.report)
        self.assertFalse(window.export_button.isEnabled())

    def test_invalid_colors_are_rejected_before_rendering(self):
        for value in ("red", "#123", "#GGGGGG", "#12345678", "123456", None):
            with self.subTest(value=value), self.assertRaises(ValueError):
                ChartColors(bar=value)


if __name__ == "__main__":
    unittest.main()
