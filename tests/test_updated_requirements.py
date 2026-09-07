"""新需求的边界值与实际导出数据契约。"""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from dataclasses import replace
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from pptx import Presentation
from PyQt6.QtWidgets import QApplication, QMessageBox

from excel_visualizer.data_service import build_report, guess_columns, read_sheet, ExcelDataError
from excel_visualizer.issues_service import build_issues_report, guess_issue_columns, IssueMetrics
from excel_visualizer.main_window import MainWindow
from excel_visualizer.ppt_service import export_pptx
from excel_visualizer.details import detail_text


class UpdatedRequirements(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.path = self.folder / "monthly.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["客户名称", "山头", "Uptime", "跑货量", "机台编码", "其他"])
        for row in [("甲", "XRF", .98, None, "001"), ("乙", "MBI", None, 50, "002"),
                    ("丙", "BFI", "97%", 0, "003"), ("丁", "XRF", None, None, "004")]:
            sheet.append(row)
        sheet["A4"].fill = PatternFill("solid", fgColor="FFFF00")
        book.save(self.path)
        book.close()

    def report(self):
        sheet = read_sheet(self.path, "Sheet")
        return build_report(sheet, guess_columns(sheet))

    def test_missing_values_mbi_three_line_labels_and_embedded_blanks(self):
        report = self.report()
        first = report.groups[0].records[0]
        self.assertEqual(first.uptime, 98)
        self.assertIsNone(first.volume)
        self.assertEqual(first.label, "001\nXRF\n甲")
        self.assertEqual(report.groups[1].records[0].product, "MBI")
        self.assertTrue(any("第 5 行，机台 004" in note and "Uptime、Run货量" in note for note in report.notes))
        path = export_pptx(report, self.folder / "missing.pptx", "2030-04")
        ppt = Presentation(path)
        for group, slide in zip(report.groups, ppt.slides):
            chart = next(s.chart for s in slide.shapes if s.has_chart)
            book = load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob), data_only=True)
            for row, record in zip(list(book.active.values)[1:], group.records):
                self.assertEqual(row[1], None if record.uptime is None else record.uptime / 100)
                self.assertEqual(row[2], record.volume)
            book.close()
            title = slide.shapes[0].text_frame.paragraphs[0]
            self.assertIn("2030年4月Uptime&Run货量", title.text)
            self.assertEqual(title.font.size.pt, 16)
            self.assertTrue(title.font.bold)
            self.assertIn('typeface="Microsoft YaHei"', slide._element.xml)

    def test_gui_warnings_detail_fields_and_period_invalidation(self):
        window = MainWindow()
        self.addCleanup(window.close)
        self.assertFalse(window.plot_button.isEnabled())
        self.assertTrue(all(not b.isEnabled() for b in window.chart_buttons))
        with patch.object(QMessageBox, "warning") as warning:
            window.open_excel(self.path)
            window._generate_report()
        self.assertIsNotNone(window.report)
        self.assertTrue(warning.called)
        self.assertEqual(window.column_boxes['machine'].count(), 7)
        text = detail_text(window, 0, True)
        for required in ("统计目的", "数据筛选条件", "横坐标", "纵坐标", "计算公式", "涉及字段", "当前自动匹配的Excel列", "注意事项", "机台编码 [E]"):
            self.assertIn(required, text)
        window.year_box.setValue(2030)
        self.assertIsNone(window.report)
        self.assertFalse(window.export_button.isEnabled())
        self.assertEqual(window.issues_page.year, 2030)

    def test_noncurrent_year_continuous_months_and_uniform_denominators(self):
        path = self.folder / 'issues.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.append(["山头", "创建时间", "机台编码", "服务请求状态"])
        for row in [("MBI", "2030-3-1", "001", " 已结束 "), ("MBI", "2030-5-2", "001", "处理中"),
                    ("MBI", "2029-3-1", "002", "已结束"), ("XRF", "2030-12-1", None, "已取消")]:
            sheet.append(row)
        book.save(path)
        book.close()
        data = read_sheet(path, 'Sheet')
        report = build_issues_report(data, guess_issue_columns(data), year=2030, start_month=3, end_month=6)
        metrics = {m.product: m for m in report.metrics}
        self.assertEqual(report.months, (3, 4, 5, 6))
        self.assertEqual(metrics['MBI'].monthly, (1, 0, 1, 0))
        self.assertEqual(metrics['MBI'].density, 2)
        self.assertEqual(metrics['MBI'].closure_rate, .5)
        self.assertEqual(metrics['XRF'].monthly, (0, 0, 0, 0))
        self.assertIsNone(metrics['XRF'].density)
        self.assertIsNone(IssueMetrics('empty', 0, (), 0, 0, 0).closure_rate)
        for start, end in [(7, 6), (7, 8)]:
            with self.assertRaises(ExcelDataError):
                build_issues_report(data, guess_issue_columns(data), year=2030, start_month=start, end_month=end)

    def test_all_empty_classification_blocks_export(self):
        report = self.report()
        empty = replace(report.groups[1], records=tuple(replace(r, volume=None) for r in report.groups[1].records))
        report = replace(report, groups=(report.groups[0], empty, report.groups[2]))
        with self.assertRaisesRegex(ValueError, '当前时间范围无有效数据'):
            export_pptx(report, self.folder / 'empty.pptx')
