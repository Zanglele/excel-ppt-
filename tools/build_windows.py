"""在当前 Windows Python 环境构建独立试用包，不覆盖或删除旧包。"""

from datetime import datetime
from importlib.metadata import version
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys


def create_samples(folder: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    folder.mkdir()
    book = Workbook()
    sheet = book.active
    sheet.title = "月报（模拟）"
    sheet.append(["客户名字", "山头", "uptime", "跑货量", "机台编码"])
    names = ("XRF", "XPS", "XRD", "AFM", "BFI", "DFI", "DBO", "IBO", "eMBI", "PC", "MBI")
    for i, name in enumerate(names):
        sheet.append([f"模拟客户{i + 1}", name, 95 + i * .3, 1000 + i * 120, f"SN-{i:03}"])
    sheet.append(["模拟未保客户", "XRF", 93, 500, "SN-uninsured"])
    sheet.cell(sheet.max_row, 1).fill = PatternFill("solid", fgColor="FFFF00")
    for column in "ABCD":
        sheet.column_dimensions[column].width = 24
    book.save(folder / "第一步模拟数据.xlsx")
    book.close()
    book = Workbook()
    sheet = book.active
    sheet.title = "问题（模拟）"
    sheet.append(["山头", "创建时间", "产品序列号/机台编码", "服务请求状态"])
    today = datetime.now()
    end_month = today.month
    for i, name in enumerate(names):
        for month in range(1, end_month + 1):
            for j in range((i + month * 2) % 8 + 1):
                sheet.append([name, f"{today.year}-{month}-{j + 1}", f"{name}-{j % 3:03}",
                              ("申请关闭", "已结束", "已取消", "处理中")[(i + j) % 4]])
    for column in "ABCD":
        sheet.column_dimensions[column].width = 30
    book.save(folder / "第二步模拟数据.xlsx")
    book.close()


def main() -> None:
    if sys.platform != "win32" or platform.machine().upper() not in ("AMD64", "X86_64"):
        raise SystemExit("请在 Windows x64 的 Python 环境中打包。")
    root = Path(__file__).resolve().parents[1]
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    release = root / "releases" / stamp
    release.mkdir(parents=True)
    versions = {name: version(name) for name in ("pyinstaller", "PyQt6", "openpyxl", "matplotlib", "python-pptx", "pandas")}
    # 只给本次打包进程提供 Python 和 Windows 路径，避免收集其他软件的同名 DLL。
    environment = {key: value for key, value in os.environ.items()
                   if not key.upper().startswith(("PYTHON", "QT_", "PYSIDE"))
                   and key.upper() not in ("MPLCONFIGDIR", "MPLBACKEND")}
    windows = Path(os.environ["SystemRoot"])
    environment["PATH"] = os.pathsep.join(map(str, (
        Path(sys.executable).parent, Path(sys.base_prefix), Path(sys.base_prefix) / "DLLs",
        windows / "System32", windows,
    )))
    environment["PYINSTALLER_CONFIG_DIR"] = str(root / "build" / stamp / "cache")
    subprocess.run([sys.executable, "-m", "PyInstaller", str(root / "packaging" / "ExcelPPT.spec"),
                    "--distpath", str(release), "--workpath", str(root / "build" / stamp)],
                   cwd=root, env=environment, check=True)
    package = release / "ExcelPPT"
    shutil.copy2(root / "packaging" / "使用说明.txt", package / "使用说明.txt")
    create_samples(package / "模拟数据")
    (release / "build-info.json").write_text(json.dumps(
        {"built_at": stamp, "python": sys.version, "packages": versions}, ensure_ascii=False, indent=2), encoding="utf-8")
    archive = shutil.make_archive(str(release / "ExcelPPT-Windows-x64"), "zip", root_dir=release, base_dir="ExcelPPT")
    print(f"\nEXE: {package / 'ExcelPPT.exe'}\nZIP: {archive}", flush=True)


if __name__ == "__main__":
    main()
