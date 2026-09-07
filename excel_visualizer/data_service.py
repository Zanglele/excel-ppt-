"""读取 Excel 的值和填充色，按月报规则校验、分组。"""

from collections import Counter
from colorsys import rgb_to_hsv
from dataclasses import dataclass
from math import isfinite
from pathlib import Path
import re
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter


NON_OPTICAL = ("XRF", "XPS", "XRD", "AFM")
OPTICAL = ("BFI", "DFI", "DBO", "IBO", "eMBI", "PC")
GROUP_TITLES = ("入质保 · 非光学产品", "入质保 · 光学产品", "未入质保 · 所有产品")
FIELDS = ("customer", "product", "uptime", "volume")
FIELD_LABELS = ("客户名字", "山头 / 产品", "Uptime", "跑货量")
ALIASES = {
    "customer": ("客户名字", "客户名称", "客户", "所有客户的名字", "customer", "customername"),
    "product": ("山头", "产品", "产品名称", "产品类型", "product", "tool", "山头/产品"),
    "uptime": ("uptime", "uptime(%)", "uptime（%）", "uptime%", "稼动率", "开机率"),
    "volume": ("跑货量", "产量", "跑货数量", "volume", "throughput"),
}


class ExcelDataError(ValueError):
    """文件或数据不符合可生成月报的条件。"""


@dataclass(frozen=True)
class CellData:
    value: object
    number_format: str
    rgb: str | None
    coordinate: str
    formula: bool = False
    merged_child: bool = False
    unsupported_fill: bool = False


@dataclass(frozen=True)
class SheetData:
    columns: tuple[str, ...]
    rows: tuple[tuple[CellData, ...], ...]
    row_numbers: tuple[int, ...]
    conditional_columns: frozenset[int]
    source: str
    sheet_name: str


@dataclass(frozen=True)
class ReportRecord:
    product: str
    customer: str
    uptime: float  # 百分制，0 到 100
    volume: float
    source_row: int

    @property
    def label(self) -> str:
        return f"{self.product}\n{self.customer}"


@dataclass(frozen=True)
class ReportGroup:
    title: str
    records: tuple[ReportRecord, ...]


@dataclass(frozen=True)
class MonthlyReport:
    groups: tuple[ReportGroup, ...]
    source: str
    sheet_name: str
    notes: tuple[str, ...] = ()


def _check_path(file_path: str | Path) -> Path:
    path = Path(file_path)
    if path.suffix.lower() != ".xlsx":
        raise ExcelDataError("请选择 .xlsx 文件；旧版 .xls 请先在 Excel 中另存为 .xlsx。")
    return path


def list_sheets(file_path: str | Path) -> list[str]:
    try:
        book = load_workbook(_check_path(file_path), read_only=True)
        try:
            return book.sheetnames
        finally:
            book.close()
    except ExcelDataError:
        raise
    except Exception as error:
        raise ExcelDataError(f"无法读取 Excel：{error}") from error


def _theme_colors(book) -> list[str]:
    if not book.loaded_theme:
        return []
    root = ElementTree.fromstring(book.loaded_theme)
    scheme = root.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme")
    if scheme is None:
        return []
    colors = {item.tag.rsplit("}", 1)[-1]: next(iter(item)).get("lastClr")
              or next(iter(item)).get("val", "") for item in scheme}
    names = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink")
    return [colors.get(name, "") for name in names]


def _fill_rgb(cell, themes: list[str], indexed_colors) -> str | None:
    if cell.fill.patternType != "solid":
        return None
    color = cell.fill.fgColor
    rgb = None
    if color.type == "rgb":
        rgb = color.rgb[-6:]
    elif color.type == "indexed" and 0 <= color.indexed < len(indexed_colors):
        rgb = indexed_colors[color.indexed][-6:]
    elif color.type == "theme" and 0 <= color.theme < len(themes):
        rgb = themes[color.theme]
    if not rgb or not re.fullmatch(r"[0-9a-fA-F]{6}", rgb):
        return None
    channels = [int(rgb[i:i + 2], 16) for i in (0, 2, 4)]
    tint = color.tint
    channels = [round(v * (1 + tint) if tint < 0 else v * (1 - tint) + 255 * tint) for v in channels]
    return "".join(f"{min(255, max(0, value)):02X}" for value in channels)


def is_yellow(rgb: str | None) -> bool:
    """覆盖标准黄、浅黄和常用金黄色，排除绿色、橙色与白色。"""
    if rgb is None:
        return False
    hue, saturation, value = rgb_to_hsv(*(int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)))
    return 42 <= hue * 360 <= 70 and saturation >= 0.15 and value >= 0.65


def read_sheet(file_path: str | Path, sheet_name: str, header_row: int = 1) -> SheetData:
    """保留颜色和数字格式；合并区域从左上角继承客户/产品信息。"""
    if header_row < 1:
        raise ExcelDataError("表头行必须大于等于 1。")
    path = _check_path(file_path)
    book = values = None
    try:
        book = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        sheet, value_sheet = book[sheet_name], values[sheet_name]
        # Worksheet 的稀疏单元格表只包含已加载的单元格；避免按格式撑大的
        # max_row/max_column 扫描整块空白区域。列宽由实际内容决定，而非表头行。
        populated = tuple(cell for cell in sheet._cells.values()
                          if cell.row >= header_row and cell.value is not None)
        if not any(cell.row == header_row for cell in populated):
            raise ExcelDataError("所选表头行为空，请选择实际列名所在的行。")
        width = max(cell.column for cell in populated)
        data_rows = sorted({cell.row for cell in populated if cell.row > header_row})
        last_row = max(cell.row for cell in populated)
        header_cells = [sheet.cell(header_row, column) for column in range(1, width + 1)]
        columns = tuple(f"{cell.value if cell.value is not None else '未命名'} [{get_column_letter(column)}]"
                        for column, cell in enumerate(header_cells, 1))
        themes = _theme_colors(book)
        indexed = getattr(book, "_colors", COLOR_INDEX)
        merged_by_row: dict[int, list] = {}
        for area in sheet.merged_cells.ranges:
            for row in range(max(header_row + 1, area.min_row), min(area.max_row, last_row) + 1):
                merged_by_row.setdefault(row, []).append(area)
        rows, row_numbers = [], []
        for number in data_rows:
            row = tuple(sheet.cell(number, column) for column in range(1, width + 1))
            if all(cell.value is None for cell in row):
                continue
            parsed = []
            for cell in row:
                source = cell
                merged_child = False
                for area in merged_by_row.get(cell.row, ()):
                    if area.min_col <= cell.column <= area.max_col:
                        source = sheet.cell(area.min_row, area.min_col)
                        merged_child = source.coordinate != cell.coordinate
                        break
                rgb = _fill_rgb(source, themes, indexed)
                parsed.append(CellData(
                    value_sheet[source.coordinate].value, source.number_format, rgb,
                    cell.coordinate, source.data_type == "f", merged_child,
                    source.fill.patternType not in (None, "solid")
                    or (source.fill.patternType == "solid" and rgb is None),
                ))
            rows.append(tuple(parsed))
            row_numbers.append(row[0].row)
        conditional_columns = set()
        for conditional in sheet.conditional_formatting:
            for area in conditional.sqref.ranges:
                if area.max_row > header_row:
                    conditional_columns.update(range(area.min_col - 1, min(area.max_col, width)))
        return SheetData(columns, tuple(rows), tuple(row_numbers), frozenset(conditional_columns),
                         str(path.resolve()), sheet_name)
    except ExcelDataError:
        raise
    except Exception as error:
        raise ExcelDataError(f"读取工作表失败：{error}") from error
    finally:
        if book is not None:
            book.close()
        if values is not None:
            values.close()


def guess_columns(sheet: SheetData) -> dict[str, int]:
    names = [re.sub(r"\s+", "", name.rsplit(" [", 1)[0]).casefold() for name in sheet.columns]
    return {field: next((i for i, name in enumerate(names) if name in
                        {alias.casefold() for alias in ALIASES[field]}), -1) for field in FIELDS}


def header_hint(sheet: SheetData, aliases: dict[str, tuple[str, ...]], guesses: dict[str, int]) -> str:
    """缺少字段时提示可能的表头位置，不擅自跳过标题或更改用户所选行。"""
    if all(index >= 0 for index in guesses.values()):
        return ""
    normalize = lambda value: re.sub(r"\s+", "", str(value)).casefold()
    field_names = [set(map(normalize, names)) for names in aliases.values()]
    candidates = []
    for number, row in zip(sheet.row_numbers[:100], sheet.rows[:100]):
        values = {normalize(cell.value) for cell in row if cell.value is not None}
        if all(values & names for names in field_names):
            candidates.append(number)
    if len(candidates) == 1:
        return f"疑似实际表头在第 {candidates[0]} 行，请修改“表头行”后点击“读取工作表 / 刷新”。"
    return "未完整识别四个字段：请确认工作表和表头行；空表头列仍按 Excel 列字母保留，可手动选择。"


def _text(cell: CellData) -> str:
    if cell.value is None or not str(cell.value).strip():
        reason = "公式没有缓存结果，请用 Excel 重新计算并保存" if cell.formula else "不能为空"
        raise ExcelDataError(f"{cell.coordinate} {reason}")
    if cell.formula and isinstance(cell.value, str) and cell.value.startswith("#"):
        raise ExcelDataError(f"{cell.coordinate} 公式错误：{cell.value}")
    return str(cell.value).strip()


def _number(cell: CellData, percentage: bool, uptime_mode: str) -> float:
    text = _text(cell).replace("，", ",").replace("％", "%")
    if cell.merged_child:
        raise ExcelDataError(f"{cell.coordinate} 数值使用了合并单元格，无法确定每条记录的值")
    if isinstance(cell.value, bool):
        raise ExcelDataError(f"{cell.coordinate} 不能使用布尔值")
    explicit_percent = text.endswith("%")
    if explicit_percent and not percentage:
        raise ExcelDataError(f"{cell.coordinate} 跑货量不能是百分比")
    numeric_text = text[:-1].strip() if explicit_percent else text
    if not re.fullmatch(r"[+-]?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", numeric_text):
        raise ExcelDataError(f"{cell.coordinate} 不是有效数值：{text}")
    number = float(numeric_text.replace(",", ""))
    number_format = re.sub(r'"[^"]*"|\\.', "", cell.number_format)
    if percentage and not explicit_percent:
        if uptime_mode == "fraction" or (uptime_mode == "excel" and "%" in number_format):
            number *= 100
    if not isfinite(number) or number < 0 or (percentage and number > 100):
        raise ExcelDataError(f"{cell.coordinate} {'Uptime 必须在 0–100% 之间' if percentage else '跑货量必须是有限非负数'}")
    return number


def build_report(sheet: SheetData, mapping: dict[str, int], uptime_mode: str = "excel") -> MonthlyReport:
    if uptime_mode not in ("excel", "points", "fraction"):
        raise ExcelDataError("未知的 Uptime 数值格式。")
    if any(mapping.get(field, -1) not in range(len(sheet.columns)) for field in FIELDS):
        raise ExcelDataError("请为客户名字、山头、Uptime、跑货量选择对应列。")
    if len({mapping[field] for field in FIELDS}) != 4:
        raise ExcelDataError("四个字段必须对应四个不同的列。")
    if mapping["customer"] in sheet.conditional_columns:
        raise ExcelDataError("客户列包含条件格式，无法可靠读取其最终显示颜色。请在 Excel 中把质保标记改为直接的黄色单元格填充后再导入。")
    canonical = {name.casefold(): name for name in NON_OPTICAL + OPTICAL}
    grouped: list[list[ReportRecord]] = [[], [], []]
    errors = []
    for row_number, row in zip(sheet.row_numbers, sheet.rows):
        cells = {field: row[mapping[field]] for field in FIELDS}
        if all(cell.value is None and not cell.formula for cell in cells.values()):
            continue
        try:
            customer, product = _text(cells["customer"]), _text(cells["product"])
            if cells["customer"].unsupported_fill:
                raise ExcelDataError(f"{cells['customer'].coordinate} 填充色无法可靠识别，请改用纯色填充")
            product = canonical.get(product.casefold(), product)
            if is_yellow(cells["customer"].rgb):
                group_index = 2
            elif product in NON_OPTICAL:
                group_index = 0
            elif product in OPTICAL:
                group_index = 1
            else:
                raise ExcelDataError(f"未标黄的产品“{product}”不在指定的 10 种产品中，请核对名称或质保标记")
            grouped[group_index].append(ReportRecord(product, customer,
                _number(cells["uptime"], True, uptime_mode),
                _number(cells["volume"], False, uptime_mode), row_number))
        except ExcelDataError as error:
            errors.append(f"第 {row_number} 行：{error}")
    if errors:
        raise ExcelDataError(f"发现 {len(errors)} 行数据问题，请修正后生成（未跳过错误行）：\n" +
                             "\n".join(errors[:20]) + ("\n其余问题请修正后重试。" if len(errors) > 20 else ""))
    if not any(grouped):
        raise ExcelDataError("没有可生成月报的数据。")
    order = {name: i for i, name in enumerate(NON_OPTICAL + OPTICAL)}
    notes = []
    for group in grouped:
        group.sort(key=lambda record: (order.get(record.product, len(order)), record.product.casefold(), record.source_row))
        repeated = Counter((r.product, r.customer) for r in group)
        if any(count > 1 for count in repeated.values()):
            notes.append("存在相同产品和客户的多条记录，已逐行保留，未自动平均 Uptime 或合计跑货量。")
        if len(group) > 30:
            notes.append("有图表超过 30 条记录，固定三页 PPT 的横轴可能较密；程序会保留全部记录。")
    return MonthlyReport(tuple(ReportGroup(title, tuple(records)) for title, records in zip(GROUP_TITLES, grouped)),
                         sheet.source, sheet.sheet_name, tuple(dict.fromkeys(notes)))
