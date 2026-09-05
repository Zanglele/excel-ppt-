"""导出前三页月报及可选的问题统计页，所有图表均可编辑。"""

from copy import deepcopy
from pathlib import Path
import os
import tempfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_DATA_LABEL_POSITION
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from excel_visualizer.data_service import MonthlyReport, ReportGroup
from excel_visualizer.chart_style import ChartColors, DEFAULT_COLORS
from excel_visualizer.issues_service import IssuesReport
from excel_visualizer.issues_charts import issue_charts, series_color


def _element(tag: str, **attributes):
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(key, str(value))
    return element


def _set_child(parent, tag: str, value) -> None:
    matches = parent.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/chart}}{tag}")
    if matches:
        matches[0].set("val", str(value))
    else:
        parent.append(_element(f"c:{tag}", val=value))


def _textbox(slide, text: str, x: float, y: float, w: float, h: float, size: int, color: str = "243247"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(color)
    return box


def _add_combo_chart(slide, group: ReportGroup, colors: ChartColors) -> None:
    volume_format = "#,##0" if all(float(r.volume).is_integer() for r in group.records) else "#,##0.00"
    data = CategoryChartData()
    data.categories = [record.label for record in group.records]
    data.add_series("Uptime（左轴）", [r.uptime / 100 for r in group.records], number_format="0.0%")
    data.add_series("跑货量（右轴）", [r.volume for r in group.records], number_format=volume_format)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.3), Inches(1.15),
                                  Inches(15.4), Inches(7.25), data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.font.name = "Microsoft YaHei"
    chart.font.size = Pt(12 if len(group.records) <= 20 else 9)
    chart.legend.font.size = Pt(14)
    category_axis = chart.category_axis
    category_axis.tick_labels.font.name = "Microsoft YaHei"
    category_axis.tick_labels.font.size = Pt(12 if len(group.records) <= 20 else 9)
    category_axis.has_title = True
    category_axis.axis_title.text_frame.text = "产品 / 客户"
    value_axis = chart.value_axis
    value_axis.minimum_scale = 0
    value_axis.maximum_scale = 1.05
    value_axis.major_unit = 0.2
    value_axis.tick_labels.number_format = "0%"
    value_axis.tick_labels.number_format_is_linked = False
    value_axis.has_title = True
    value_axis.axis_title.text_frame.text = "Uptime (%)"
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("D7DDE5")
    value_axis.major_gridlines.format.line.width = Pt(0.5)
    line = chart.series[0]
    line.format.line.color.rgb = RGBColor.from_string(colors.line[1:])
    line.format.line.width = Pt(2.5)
    line.marker.style = XL_MARKER_STYLE.CIRCLE
    line.marker.size = 5
    line.marker.format.fill.solid()
    line.marker.format.fill.fore_color.rgb = RGBColor.from_string(colors.line[1:])
    line.marker.format.line.color.rgb = RGBColor.from_string(colors.line[1:])
    volume = chart.series[1]
    volume.format.fill.solid()
    volume.format.fill.fore_color.rgb = RGBColor.from_string(colors.bar[1:])
    volume.format.line.fill.background()

    # python-pptx 尚无创建组合图的公开接口；只在此函数内补全 OOXML 的第二绘图区和轴。
    plot_area = chart._chartSpace.chart.plotArea
    line_plot = plot_area.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}lineChart")
    volume_series = line_plot.findall("{http://schemas.openxmlformats.org/drawingml/2006/chart}ser")[1]
    line_plot.remove(volume_series)
    for tag in ("marker", "smooth"):
        for child in volume_series.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/chart}}{tag}"):
            volume_series.remove(child)
    primary_cat = plot_area.xpath("./c:catAx")[0]
    primary_val = plot_area.xpath("./c:valAx")[0]
    secondary_cat, secondary_val = deepcopy(primary_cat), deepcopy(primary_val)
    category_id, value_id = "100100", "100101"
    bar_plot = _element("c:barChart")
    bar_plot.append(_element("c:barDir", val="col"))
    bar_plot.append(_element("c:grouping", val="clustered"))
    bar_plot.append(volume_series)
    bar_plot.append(_element("c:gapWidth", val="85"))
    bar_plot.append(_element("c:axId", val=category_id))
    bar_plot.append(_element("c:axId", val=value_id))
    plot_area.insert(list(plot_area).index(line_plot), bar_plot)
    _set_child(secondary_cat, "axId", category_id)
    _set_child(secondary_cat, "crossAx", value_id)
    _set_child(secondary_cat, "delete", 1)
    # 只保留一条横轴标题。
    for title in secondary_cat.xpath("./c:title"):
        secondary_cat.remove(title)
    _set_child(secondary_val, "axId", value_id)
    _set_child(secondary_val, "axPos", "r")
    _set_child(secondary_val, "crossAx", category_id)
    _set_child(secondary_val, "crosses", "max")
    for grid in secondary_val.xpath("./c:majorGridlines"):
        secondary_val.remove(grid)
    for unit in secondary_val.xpath("./c:majorUnit"):
        secondary_val.remove(unit)
    for text in secondary_val.xpath("./c:title//a:t"):
        text.text = "跑货量"
    secondary_val.xpath("./c:numFmt")[0].set("formatCode", volume_format)
    scaling = secondary_val.xpath("./c:scaling")[0]
    scaling.xpath("./c:max")[0].set("val", str(max(1, max(r.volume for r in group.records) * 1.18)))
    for axis in (primary_val, secondary_val):
        _set_child(axis, "crossBetween", "between")
    # 每条记录均保留一个横轴标签，禁止自动隔条省略。
    _set_child(primary_cat, "tickLblSkip", 1)
    _set_child(primary_cat, "tickMarkSkip", 1)
    if len(group.records) > 12:
        for body in primary_cat.xpath("./c:txPr/a:bodyPr"):
            body.set("rot", "-2700000")
    plot_area.append(secondary_cat)
    plot_area.append(secondary_val)


def _add_issues_slide(presentation, report: IssuesReport) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(slide, "问题统计", 0.4, 0.06, 8, 0.48, 22)
    _textbox(slide, f"月度趋势截至 {report.as_of:%Y-%m}", 11.7, 0.12, 3.8, 0.35, 11, "667085")
    for index, spec in enumerate(issue_charts(report)):
        data = CategoryChartData()
        # 空比值保持为空，横轴标注 N/A，不向图表写入虚假的零值。
        data.categories = [name + ("\nN/A" if spec.kind == "bar" and spec.series[0][1][i] is None else "")
                           for i, name in enumerate(spec.categories)]
        for name, values in spec.series:
            data.add_series(name, values, number_format=spec.number_format)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE_MARKERS if spec.kind == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(0.3 + (index % 2) * 7.9), Inches(0.58 + (index // 2) * 1.95),
            Inches(7.5), Inches(1.9), data).chart
        chart.font.name = "Microsoft YaHei"
        chart.font.size = Pt(8)
        chart.has_title = True
        chart.chart_title.text_frame.text = spec.title
        for paragraph in chart.chart_title.text_frame.paragraphs:
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(12)
            paragraph.font.bold = True
        chart.has_legend = spec.kind == "line"
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(7)
        value_axis = chart.value_axis
        value_axis.minimum_scale = 0
        value_axis.tick_labels.number_format = spec.number_format
        value_axis.tick_labels.number_format_is_linked = False
        value_axis.has_title = True
        value_axis.axis_title.text_frame.text = spec.ylabel
        for paragraph in value_axis.axis_title.text_frame.paragraphs:
            paragraph.font.size = Pt(8)
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("E4E7EC")
        maximum = max((v for _, values in spec.series for v in values if v is not None), default=0)
        value_axis.maximum_scale = 1.15 if spec.number_format.endswith("%") else max(1, maximum * 1.2)
        if spec.number_format == "0" and maximum < 5:
            value_axis.major_unit = 1
        category_axis = chart.category_axis
        category_axis.tick_labels.font.size = Pt(8)
        _set_child(category_axis._element, "tickLblSkip", 1)
        _set_child(category_axis._element, "tickMarkSkip", 1)
        for i, series in enumerate(chart.series):
            color = RGBColor.from_string(series_color(i if spec.kind == "line" else 0)[1:])
            if spec.kind == "line":
                series.format.line.color.rgb = color
                series.format.line.width = Pt(1.3)
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.size = 3
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
                series.marker.format.line.color.rgb = color
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
                series.format.line.fill.background()
        if spec.kind == "bar":
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.number_format = spec.number_format
            plot.data_labels.font.size = Pt(7)
    # 页面下半部不添加页脚、说明或占位文字，保留给总结性批注。
    slide.notes_slide.notes_text_frame.text = (
        f"数据来源：{Path(report.source).name}；工作表：{report.sheet_name}\n"
        f"前两图为2026年；月度趋势范围：1—{report.months[-1]}月。后两图统计范围：{report.scope}。\n"
        "每行计一次问题（重复行也计数）；机器编号在各山头内去重，空编号不计机器。\n"
        "关闭率 =（申请关闭 + 已结束 + 已取消）/ 总问题数。N/A 表示分母为零。\n"
        + "\n".join(report.notes)
    )


def export_issues_pptx(report: IssuesReport, output_path: str | Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(16)
    presentation.slide_height = Inches(9)
    _add_issues_slide(presentation, report)
    return _save_presentation(presentation, output_path)


def export_pptx(report: MonthlyReport, output_path: str | Path, period: str = "",
                colors: ChartColors = DEFAULT_COLORS, *, issues_report: IssuesReport | None = None) -> Path:
    """先写临时文件再替换，失败时保留已有 PPT。"""
    if len(report.groups) != 3:
        raise ValueError("月报必须包含固定的三个分组。")
    presentation = Presentation()
    presentation.slide_width = Inches(16)
    presentation.slide_height = Inches(9)
    for index, group in enumerate(report.groups, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _textbox(slide, group.title, 0.55, 0.25, 12.5, 0.65, 30)
        _textbox(slide, period.strip(), 13.1, 0.35, 2.3, 0.5, 16, "667085")
        if group.records:
            _add_combo_chart(slide, group, colors)
        else:
            _textbox(slide, "本月无符合条件的数据", 4.8, 4, 7, 1, 24, "667085")
        _textbox(slide, f"共 {len(group.records)} 条记录", 0.6, 8.5, 8, 0.35, 12, "667085")
        _textbox(slide, f"{index} / {4 if issues_report else 3}", 14.5, 8.5, 1, 0.35, 12, "667085")
        slide.notes_slide.notes_text_frame.text = (
            f"数据来源：{Path(report.source).name}；工作表：{report.sheet_name}\n"
            "黄色客户单元格代表未入质保；其他客户按产品名单归类。\n"
            + "\n".join(f"Excel 第 {r.source_row} 行：{r.product} / {r.customer}" for r in group.records)
        )
    if issues_report is not None:
        _add_issues_slide(presentation, issues_report)
    return _save_presentation(presentation, output_path)


def _save_presentation(presentation, output_path: str | Path) -> Path:
    path = Path(output_path)
    if path.suffix.lower() != ".pptx":
        path = path.with_suffix(".pptx")
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".pptx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        presentation.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()
    return path
