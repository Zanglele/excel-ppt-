"""显式 --self-test 模式：验证冻结程序的 Qt、Excel 和 PPT 完整链路。"""

import json
from pathlib import Path
import sys
import traceback


def run_check(output_directory: str) -> int:
    folder = Path(output_directory).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    result = {"frozen": bool(getattr(sys, "frozen", False)), "executable": sys.executable}
    window = None
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import PatternFill
        from pptx import Presentation
        from PyQt6.QtWidgets import QApplication, QFileDialog, QMessageBox
        from excel_visualizer.main_window import MainWindow

        def fail_dialog(*args, **kwargs):
            raise RuntimeError(str(args[1:]))

        originals = (QMessageBox.warning, QMessageBox.critical, QMessageBox.information,
                     QFileDialog.getSaveFileName)
        QMessageBox.warning = staticmethod(fail_dialog)
        QMessageBox.critical = staticmethod(fail_dialog)
        QMessageBox.information = staticmethod(lambda *args, **kwargs: None)
        try:
            application = QApplication.instance() or QApplication([])
            result["qt_platform"] = application.platformName()
            window = MainWindow()
            window.year_box.setValue(2026)
            window.month_box.setValue(9)
            window.show()
            application.processEvents()
            first = folder / "第一步模拟.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["客户名字", "山头", "uptime", "跑货量", "机台编码"])
            sheet.append(["模拟客户甲", "XRF", 98, 1234, "001"])
            sheet.append(["模拟客户乙", "BFI", 97, 567, "002"])
            sheet.append(["模拟未保客户", "XRF", 96, 89, "003"])
            sheet["A4"].fill = PatternFill("solid", fgColor="FFFF00")
            book.save(first)
            book.close()
            second = folder / "第二步模拟.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["虚构导出表", "标题行只有两个单元格"])
            sheet.merge_cells("B1:H1")
            sheet.append(["实际表头在第3行"])
            sheet.append(["编号", "说明", "山头", "创建时间", "其他列", "产品序列号/机台编码", "地区", "服务请求状态"])
            sheet.append([1, "模拟", "XRF", "2026-1-2", "模拟", "001", "模拟", "申请关闭"])
            sheet.append([2, "模拟", "XRF", "2026-1-3", "模拟", "001", "模拟", "已结束"])
            sheet.append([3, "模拟", "XRF", "2026-2-4", "模拟", "002", "模拟", "处理中"])
            sheet.append([4, "模拟", "BFI", "2026-1-2", "模拟", "003", "模拟", "已取消"])
            book.save(second)
            book.close()
            window.open_excel(first)
            window._generate_report()
            assert window.report is not None
            assert [len(group.records) for group in window.report.groups] == [1, 1, 1]
            application.processEvents()
            window.workflow_tabs.setCurrentIndex(1)
            application.processEvents()
            window.issues_page.open_excel(second)
            assert len(window.issues_page.sheet.columns) == 8
            assert window.issues_page.column_boxes["status"].count() == 9
            assert "第 3 行" in window.issues_page.status_label.text()
            window.issues_page.header_row.setValue(3)
            window.issues_page._load_sheet()
            assert window.issues_page.column_boxes["status"].currentData() == 7
            result["header_columns"] = 8
            result["header_row"] = 3
            window.issues_page._generate_report()
            assert window.issues_page.report is not None
            xrf = window.issues_page.report.metrics[0]
            assert (xrf.annual_total, xrf.machines, xrf.closed, xrf.density) == (3, 2, 2, 1.5)
            assert window.combined_export_button.isEnabled()
            application.processEvents()
            window.grab().save(str(folder / "打包程序界面.png"))
            result["exports"] = {}
            for filename, export, count in (
                ("第一步.pptx", window._export_report, 3),
                ("第二步.pptx", window.issues_page._export_report, 1),
                ("合并报告.pptx", window._export_combined_report, 4),
            ):
                path = folder / filename
                QFileDialog.getSaveFileName = staticmethod(lambda *args, p=str(path), **kwargs: (p, ""))
                export()
                ppt = Presentation(path)
                assert len(ppt.slides) == count
                if count in (1, 4):
                    issue_slide = ppt.slides[-1]
                    assert sum(shape.has_chart for shape in issue_slide.shapes) == 4
                    assert all(shape.top + shape.height <= ppt.slide_height / 2 for shape in issue_slide.shapes)
                # 读取每张内嵌 Excel，确认模板和嵌入资源都已打包。
                from io import BytesIO
                for slide in ppt.slides:
                    for shape in slide.shapes:
                        if shape.has_chart:
                            embedded = shape.chart.part.chart_workbook.xlsx_part.blob
                            workbook = load_workbook(BytesIO(embedded), data_only=True)
                            assert workbook.active.max_row > 1
                            workbook.close()
                result["exports"][filename] = count
            result["status"] = "passed"
        finally:
            QMessageBox.warning, QMessageBox.critical, QMessageBox.information, QFileDialog.getSaveFileName = originals
    except Exception:
        result["status"] = "failed"
        result["error"] = traceback.format_exc()
    finally:
        if window is not None:
            window.close()
        (folder / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if result["status"] == "passed" else 1
